"""LocalData MCP - Database connection and query management."""

import atexit
import configparser
import hashlib
import json
import logging
import os
import psutil
import tempfile
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

import pandas as pd
import yaml
from fastmcp import FastMCP
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.sql import quoted_name

# TOML support
try:
    import toml
    TOML_AVAILABLE = True
except ImportError:
    TOML_AVAILABLE = False

# Excel support libraries with graceful error handling
try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    XLRD_AVAILABLE = False

try:
    import defusedxml
    DEFUSEDXML_AVAILABLE = True
except ImportError:
    DEFUSEDXML_AVAILABLE = False

# ODS (LibreOffice Calc) support libraries with graceful error handling
try:
    from odf import opendocument, table
    ODFPY_AVAILABLE = True
except ImportError:
    ODFPY_AVAILABLE = False

# XML parsing support
try:
    import lxml
    LXML_AVAILABLE = True
except ImportError:
    LXML_AVAILABLE = False

# Analytical format support (Parquet, Arrow, Feather)
try:
    import pyarrow
    PYARROW_AVAILABLE = True
except ImportError:
    PYARROW_AVAILABLE = False

# Apple Numbers support
try:
    from numbers_parser import Document
    NUMBERS_PARSER_AVAILABLE = True
except ImportError:
    NUMBERS_PARSER_AVAILABLE = False

# DuckDB support
try:
    import duckdb
    DUCKDB_AVAILABLE = True
except ImportError:
    DUCKDB_AVAILABLE = False

# HDF5 support for scientific data
try:
    import h5py
    H5PY_AVAILABLE = True
except ImportError:
    H5PY_AVAILABLE = False

# Modern database support - Redis
try:
    import redis
    REDIS_AVAILABLE = True
except ImportError:
    REDIS_AVAILABLE = False

# Modern database support - Elasticsearch
try:
    from elasticsearch import Elasticsearch
    ELASTICSEARCH_AVAILABLE = True
except ImportError:
    ELASTICSEARCH_AVAILABLE = False

# Modern database support - MongoDB
try:
    import pymongo
    PYMONGO_AVAILABLE = True
except ImportError:
    PYMONGO_AVAILABLE = False

# Modern database support - InfluxDB
try:
    from influxdb_client import InfluxDBClient
    INFLUXDB_AVAILABLE = True
except ImportError:
    INFLUXDB_AVAILABLE = False

# Modern database support - Neo4j
try:
    from neo4j import GraphDatabase
    NEO4J_AVAILABLE = True
except ImportError:
    NEO4J_AVAILABLE = False

# Modern database support - CouchDB
try:
    import couchdb
    COUCHDB_AVAILABLE = True
except ImportError:
    COUCHDB_AVAILABLE = False

# Set up logging
logger = logging.getLogger(__name__)
logger.setLevel(logging.INFO)

# Create console handler if none exists
if not logger.handlers:
    console_handler = logging.StreamHandler()
    console_handler.setLevel(logging.INFO)
    formatter = logging.Formatter(
        "%(asctime)s - %(name)s - %(levelname)s - %(message)s"
    )
    console_handler.setFormatter(formatter)
    logger.addHandler(console_handler)

# Create the MCP server instance
mcp = FastMCP("localdata-mcp")


@dataclass
class QueryBuffer:
    query_id: str
    db_name: str
    query: str
    results: pd.DataFrame
    timestamp: float
    source_file_path: Optional[str] = None
    source_file_mtime: Optional[float] = None


class DatabaseManager:
    def __init__(self):
        self.connections: Dict[str, Any] = {}
        self.db_types: Dict[str, str] = {}  # Track database type for each connection
        self.query_history: Dict[str, List[str]] = {}

        # Security and connection management
        self.connection_semaphore = threading.Semaphore(
            10
        )  # Max 10 concurrent connections
        self.connection_lock = threading.Lock()
        self.connection_count = 0

        # Query buffering system
        self.query_buffers: Dict[str, QueryBuffer] = {}
        self.query_buffer_lock = threading.Lock()

        # Temporary file management
        self.temp_files: List[str] = []
        self.temp_file_lock = threading.Lock()

        # Auto-cleanup for buffers (10 minute expiry)
        self.buffer_cleanup_interval = 600  # 10 minutes
        self.last_cleanup = time.time()

        # Register cleanup on exit
        atexit.register(self._cleanup_all)

    def _get_connection(self, name: str):
        if name not in self.connections:
            raise ValueError(
                f"Database '{name}' is not connected. Use 'connect_database' first."
            )
        return self.connections[name]

    def _sanitize_path(self, file_path: str):
        """Enhanced path security - restrict to current working directory and subdirectories only."""
        base_dir = Path(os.getcwd()).resolve()
        try:
            # Resolve the path to handle symlinks and relative paths
            abs_file_path = Path(file_path).resolve()
        except (OSError, ValueError) as e:
            raise ValueError(f"Invalid path '{file_path}': {e}")

        # Security check: ensure path is within base directory
        try:
            abs_file_path.relative_to(base_dir)
        except ValueError:
            raise ValueError(
                f"Path '{file_path}' is outside the allowed directory. Only current directory and subdirectories are allowed."
            )

        # Check if file exists
        if not abs_file_path.is_file():
            raise ValueError(f"File not found at path '{file_path}'.")

        return str(abs_file_path)

    def _serialize_complex_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        """Serialize complex nested objects (dict, list) in DataFrame columns to JSON strings.
        
        This prevents SQLite insertion errors when DataFrame contains nested structures
        that can't be directly serialized to database columns.
        
        Args:
            df: DataFrame that may contain columns with complex nested objects
            
        Returns:
            DataFrame with complex objects serialized as JSON strings
        """
        if df.empty:
            return df
        
        # Create a copy to avoid modifying the original DataFrame
        df_copy = df.copy()
        
        for col in df_copy.columns:
            if df_copy[col].dtype == 'object':
                # Check if this column contains complex objects (dict/list)
                non_null_values = df_copy[col].dropna()
                if not non_null_values.empty:
                    # Check first few non-null values to determine if serialization is needed
                    sample_values = non_null_values.head(3)
                    needs_serialization = any(
                        isinstance(val, (dict, list)) for val in sample_values.values
                    )
                    
                    if needs_serialization:
                        logger.info(f"Serializing column '{col}' containing complex nested objects")
                        def serialize_value(x):
                            # Handle null values safely - avoid pd.notna() with arrays
                            if x is None:
                                return x
                            # Check if it's a complex object that needs serialization
                            if isinstance(x, (dict, list)):
                                return json.dumps(x)
                            return x
                        
                        df_copy[col] = df_copy[col].apply(serialize_value)
        
        return df_copy

    def _get_file_size(self, file_path: str) -> int:
        """Get file size in bytes."""
        try:
            return os.path.getsize(file_path)
        except OSError as e:
            raise ValueError(f"Cannot get size of file '{file_path}': {e}")

    def _is_large_file(self, file_path: str, threshold_mb: int = 100) -> bool:
        """Check if file exceeds the size threshold (default 100MB)."""
        threshold_bytes = threshold_mb * 1024 * 1024
        return self._get_file_size(file_path) > threshold_bytes

    def _sanitize_sheet_name(self, sheet_name: str, used_names: set = None) -> str:
        """Sanitize sheet name for use as SQL table name."""
        import re
        
        if used_names is None:
            used_names = set()
        
        # Convert to string and strip whitespace
        name = str(sheet_name).strip()
        
        # Replace spaces and hyphens with underscores
        name = re.sub(r'[\s\-]+', '_', name)
        
        # Remove or replace problematic characters, keep only alphanumeric and underscore
        name = re.sub(r'[^\w]', '_', name)
        
        # Remove consecutive underscores
        name = re.sub(r'_+', '_', name)
        
        # Ensure it starts with letter or underscore
        if name and not re.match(r'^[a-zA-Z_]', name):
            name = 'sheet_' + name
        
        # Handle empty names
        if not name:
            name = 'sheet_unnamed'
        
        # Ensure uniqueness
        original_name = name
        counter = 1
        while name.lower() in {n.lower() for n in used_names}:
            name = f"{original_name}_{counter}"
            counter += 1
        
        # Add to used names
        used_names.add(name)
        
        return name

    def _generate_query_id(self, db_name: str, query: str) -> str:
        """Generate a unique query ID in format: {db}_{timestamp}_{4char_hash}."""
        timestamp = int(time.time())
        query_hash = hashlib.md5(query.encode(), usedforsecurity=False).hexdigest()[
            :4
        ]  # nosec B324
        return f"{db_name}_{timestamp}_{query_hash}"

    def _cleanup_expired_buffers(self):
        """Remove expired query buffers (older than 10 minutes)."""
        current_time = time.time()
        if current_time - self.last_cleanup < self.buffer_cleanup_interval:
            return  # Skip if not time for cleanup yet

        with self.query_buffer_lock:
            expired_ids = [
                query_id
                for query_id, buffer in self.query_buffers.items()
                if current_time - buffer.timestamp > self.buffer_cleanup_interval
            ]
            for query_id in expired_ids:
                del self.query_buffers[query_id]

        self.last_cleanup = current_time

    def _cleanup_all(self):
        """Clean up all resources on exit."""
        # Clean up temporary files
        with self.temp_file_lock:
            for temp_file in self.temp_files:
                try:
                    if os.path.exists(temp_file):
                        os.unlink(temp_file)
                except OSError:
                    pass  # Ignore errors during cleanup
            self.temp_files.clear()

        # Close all database connections
        with self.connection_lock:
            for name, engine in self.connections.items():
                try:
                    engine.dispose()
                except:
                    pass  # Ignore errors during cleanup
            self.connections.clear()
            self.db_types.clear()

    def _get_engine(self, db_type: str, conn_string: str, sheet_name: Optional[str] = None):
        if db_type == "sqlite":
            return create_engine(f"sqlite:///{conn_string}")
        elif db_type == "postgresql":
            return create_engine(conn_string)
        elif db_type == "mysql":
            return create_engine(conn_string)
        elif db_type == "duckdb":
            if not DUCKDB_AVAILABLE:
                raise ValueError("duckdb library is required for DuckDB connections. Install with: pip install duckdb")
            return create_engine(f"duckdb:///{conn_string}")
        elif db_type == "redis":
            if not REDIS_AVAILABLE:
                raise ValueError("redis library is required for Redis connections. Install with: pip install redis")
            return self._create_redis_connection(conn_string)
        elif db_type == "elasticsearch":
            if not ELASTICSEARCH_AVAILABLE:
                raise ValueError("elasticsearch library is required for Elasticsearch connections. Install with: pip install elasticsearch")
            return self._create_elasticsearch_connection(conn_string)
        elif db_type == "mongodb":
            if not PYMONGO_AVAILABLE:
                raise ValueError("pymongo library is required for MongoDB connections. Install with: pip install pymongo")
            return self._create_mongodb_connection(conn_string)
        elif db_type == "influxdb":
            if not INFLUXDB_AVAILABLE:
                raise ValueError("influxdb-client library is required for InfluxDB connections. Install with: pip install influxdb-client")
            return self._create_influxdb_connection(conn_string)
        elif db_type == "neo4j":
            if not NEO4J_AVAILABLE:
                raise ValueError("neo4j library is required for Neo4j connections. Install with: pip install neo4j")
            return self._create_neo4j_connection(conn_string)
        elif db_type == "couchdb":
            if not COUCHDB_AVAILABLE:
                raise ValueError("couchdb library is required for CouchDB connections. Install with: pip install couchdb")
            return self._create_couchdb_connection(conn_string)
        elif db_type in ["csv", "json", "yaml", "toml", "excel", "ods", "numbers", "xml", "ini", "tsv", "parquet", "feather", "arrow", "hdf5"]:
            sanitized_path = self._sanitize_path(conn_string)
            return self._create_engine_from_file(sanitized_path, db_type, sheet_name)
        else:
            raise ValueError(f"Unsupported db_type: {db_type}")

    def _create_engine_from_file(self, file_path: str, file_type: str, sheet_name: Optional[str] = None):
        """Create SQLite engine from file, using temporary storage for large files.
        
        Args:
            file_path: Path to the file
            file_type: Type of file (excel, ods, numbers, hdf5, etc.)
            sheet_name: For Excel/ODS/Numbers files, specific sheet to load; for HDF5 files, specific dataset to load. If None, load all sheets/datasets.
        """
        try:
            # Check if file is large
            is_large = self._is_large_file(file_path)

            # Load data based on file type
            if file_type == "csv":
                try:
                    data = pd.read_csv(file_path)
                except pd.errors.ParserError:
                    # Fallback for CSV with no header
                    data = pd.read_csv(file_path, header=None)
            elif file_type == "json":
                data = pd.read_json(file_path)
            elif file_type == "yaml":
                with open(file_path, "r") as f:
                    yaml_data = yaml.safe_load(f)
                data = (
                    pd.json_normalize(yaml_data)
                    if isinstance(yaml_data, (list, dict))
                    else pd.DataFrame(yaml_data)
                )
            elif file_type == "toml":
                if not TOML_AVAILABLE:
                    raise ValueError(
                        "toml library is required for TOML files. "
                        "Install with: pip install toml"
                    )
                with open(file_path, "r") as f:
                    toml_data = toml.load(f)
                data = (
                    pd.json_normalize(toml_data)
                    if isinstance(toml_data, (list, dict))
                    else pd.DataFrame(toml_data)
                )
            elif file_type == "excel":
                data = self._load_excel_file(file_path, sheet_name)
            elif file_type == "ods":
                data = self._load_ods_file(file_path, sheet_name)
            elif file_type == "numbers":
                data = self._load_numbers_file(file_path, sheet_name)
            elif file_type == "xml":
                data = self._load_xml_file(file_path)
            elif file_type == "ini":
                data = self._load_ini_file(file_path)
            elif file_type == "tsv":
                try:
                    data = pd.read_csv(file_path, sep='\t')
                except pd.errors.ParserError:
                    # Fallback for TSV with no header
                    data = pd.read_csv(file_path, sep='\t', header=None)
            elif file_type == "parquet":
                if not PYARROW_AVAILABLE:
                    raise ValueError(
                        "pyarrow library is required for Parquet files. "
                        "Install with: pip install pyarrow"
                    )
                data = pd.read_parquet(file_path, engine='pyarrow')
            elif file_type == "feather":
                if not PYARROW_AVAILABLE:
                    raise ValueError(
                        "pyarrow library is required for Feather files. "
                        "Install with: pip install pyarrow"
                    )
                data = pd.read_feather(file_path)
            elif file_type == "arrow":
                if not PYARROW_AVAILABLE:
                    raise ValueError(
                        "pyarrow library is required for Arrow files. "
                        "Install with: pip install pyarrow"
                    )
                data = pd.read_feather(file_path)  # Arrow IPC format uses same reader
            elif file_type == "hdf5":
                data = self._load_hdf5_file(file_path, sheet_name)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")

            # Create engine - use temporary file for large files, memory for small ones
            if is_large:
                # Create temporary SQLite file
                temp_fd, temp_path = tempfile.mkstemp(
                    suffix=".sqlite", prefix="db_client_"
                )
                os.close(
                    temp_fd
                )  # Close the file descriptor, SQLAlchemy will handle the file

                with self.temp_file_lock:
                    self.temp_files.append(temp_path)

                engine = create_engine(f"sqlite:///{temp_path}")
            else:
                engine = create_engine("sqlite:///:memory:")

            # Load data into SQLite
            # Handle multi-sheet files (Excel/ODS) vs single DataFrame files
            if isinstance(data, dict):
                # Multi-sheet file: create separate table for each sheet
                for table_name, df in data.items():
                    # Serialize complex nested objects before SQLite insertion
                    df_serialized = self._serialize_complex_columns(df)
                    df_serialized.to_sql(table_name, engine, index=False, if_exists="replace")
                    logger.info(f"Created table '{table_name}' with {len(df_serialized)} rows")
            else:
                # Single DataFrame: create single table called 'data_table'
                # Serialize complex nested objects before SQLite insertion
                data_serialized = self._serialize_complex_columns(data)
                data_serialized.to_sql("data_table", engine, index=False, if_exists="replace")
                logger.info(f"Created table 'data_table' with {len(data_serialized)} rows")
            
            return engine

        except Exception as e:
            raise ValueError(
                f"Failed to create engine from {file_type} file '{file_path}': {e}"
            )

    def _load_excel_file(self, file_path: str, sheet_name: Optional[str] = None) -> Dict[str, pd.DataFrame]:
        """Load Excel file (.xlsx or .xls) into dict of pandas DataFrames (one per sheet).
        
        Args:
            file_path: Path to the Excel file
            sheet_name: Specific sheet to load. If None, load all sheets.
        """
        # Check for security library
        if not DEFUSEDXML_AVAILABLE:
            logger.warning("defusedxml not available - Excel files may be vulnerable to XML attacks")
        
        # Detect file format based on extension
        file_ext = Path(file_path).suffix.lower()
        
        try:
            # Determine engine based on file extension
            if file_ext in ['.xlsx', '.xlsm']:
                # Modern Excel format
                if not OPENPYXL_AVAILABLE:
                    raise ValueError(
                        "openpyxl library is required for .xlsx files. "
                        "Install with: pip install openpyxl"
                    )
                engine = 'openpyxl'
                
            elif file_ext == '.xls':
                # Legacy Excel format
                if not XLRD_AVAILABLE:
                    raise ValueError(
                        "xlrd library is required for .xls files. "
                        "Install with: pip install xlrd"
                    )
                engine = 'xlrd'
                
            else:
                # Try pandas auto-detection as fallback
                logger.info(f"Unknown Excel extension '{file_ext}', trying pandas auto-detection")
                engine = None
            
            # Read all sheets or specific sheet
            with pd.ExcelFile(file_path, engine=engine) as excel_file:
                available_sheet_names = excel_file.sheet_names
                logger.info(f"Found {len(available_sheet_names)} sheets in Excel file: {available_sheet_names}")
                
                # Determine which sheets to load
                if sheet_name is not None:
                    if sheet_name not in available_sheet_names:
                        raise ValueError(f"Sheet '{sheet_name}' not found in Excel file. Available sheets: {available_sheet_names}")
                    sheets_to_load = [sheet_name]
                    logger.info(f"Loading specific sheet: {sheet_name}")
                else:
                    sheets_to_load = available_sheet_names
                    logger.info(f"Loading all sheets")
                
                # Track used table names for uniqueness
                used_names = set()
                sheets_data = {}
                
                for current_sheet_name in sheets_to_load:
                    try:
                        # Load sheet data
                        df = pd.read_excel(excel_file, sheet_name=current_sheet_name)
                        
                        # Skip empty sheets
                        if df.empty:
                            logger.warning(f"Sheet '{current_sheet_name}' is empty, skipping")
                            continue
                        
                        # Clean up column names (remove extra whitespace, replace problematic characters)
                        df.columns = [str(col).strip().replace(' ', '_').replace('-', '_') for col in df.columns]
                        
                        # Handle Excel-specific data types
                        # Convert any datetime-like columns properly
                        for col in df.columns:
                            if df[col].dtype == 'object':
                                # Try to convert to datetime if it looks like dates
                                try:
                                    pd.to_datetime(df[col], errors='raise')
                                    df[col] = pd.to_datetime(df[col], errors='coerce')
                                except:
                                    pass  # Keep as object type
                        
                        # Sanitize sheet name for SQL table name
                        table_name = self._sanitize_sheet_name(current_sheet_name, used_names)
                        sheets_data[table_name] = df
                        
                        logger.info(f"Successfully loaded sheet '{current_sheet_name}' as table '{table_name}' with {len(df)} rows and {len(df.columns)} columns")
                        
                    except Exception as e:
                        logger.warning(f"Failed to load sheet '{current_sheet_name}': {e}")
                        continue
                
                if not sheets_data:
                    raise ValueError(f"Excel file '{file_path}' contains no readable data in any sheets")
                
                logger.info(f"Successfully loaded Excel file '{file_path}' with {len(sheets_data)} sheets")
                return sheets_data
            
        except Exception as e:
            raise ValueError(f"Failed to load Excel file '{file_path}': {e}")

    def _load_ods_file(self, file_path: str, sheet_name: Optional[str] = None) -> Dict[str, pd.DataFrame]:
        """Load LibreOffice Calc ODS file into dict of pandas DataFrames (one per sheet).
        
        Args:
            file_path: Path to the ODS file
            sheet_name: Specific sheet to load. If None, load all sheets.
        """
        # Check for ODS support
        if not ODFPY_AVAILABLE:
            logger.warning("odfpy not available - ODS files cannot be loaded")
            raise ValueError(
                "odfpy library is required for .ods files. "
                "Install with: pip install odfpy"
            )
        
        try:
            # Use pandas native ODS support with odfpy engine
            with pd.ExcelFile(file_path, engine='odf') as excel_file:
                available_sheet_names = excel_file.sheet_names
                logger.info(f"Found {len(available_sheet_names)} sheets in ODS file: {available_sheet_names}")
                
                # Determine which sheets to load
                if sheet_name is not None:
                    if sheet_name not in available_sheet_names:
                        raise ValueError(f"Sheet '{sheet_name}' not found in ODS file. Available sheets: {available_sheet_names}")
                    sheets_to_load = [sheet_name]
                    logger.info(f"Loading specific sheet: {sheet_name}")
                else:
                    sheets_to_load = available_sheet_names
                    logger.info(f"Loading all sheets")
                
                # Track used table names for uniqueness
                used_names = set()
                sheets_data = {}
                
                for current_sheet_name in sheets_to_load:
                    try:
                        # Load sheet data
                        df = pd.read_excel(excel_file, sheet_name=current_sheet_name, engine='odf')
                        
                        # Skip empty sheets
                        if df.empty:
                            logger.warning(f"Sheet '{current_sheet_name}' is empty, skipping")
                            continue
                        
                        # Clean up column names (remove extra whitespace, replace problematic characters)
                        df.columns = [str(col).strip().replace(' ', '_').replace('-', '_') for col in df.columns]
                        
                        # Handle ODS-specific data types
                        # Convert any datetime-like columns properly
                        for col in df.columns:
                            if df[col].dtype == 'object':
                                # Try to convert to datetime if it looks like dates
                                try:
                                    pd.to_datetime(df[col], errors='raise')
                                    df[col] = pd.to_datetime(df[col], errors='coerce')
                                except:
                                    pass  # Keep as object type
                        
                        # Sanitize sheet name for SQL table name
                        table_name = self._sanitize_sheet_name(current_sheet_name, used_names)
                        sheets_data[table_name] = df
                        
                        logger.info(f"Successfully loaded sheet '{current_sheet_name}' as table '{table_name}' with {len(df)} rows and {len(df.columns)} columns")
                        
                    except Exception as e:
                        logger.warning(f"Failed to load sheet '{current_sheet_name}': {e}")
                        continue
                
                if not sheets_data:
                    raise ValueError(f"ODS file '{file_path}' contains no readable data in any sheets")
                
                logger.info(f"Successfully loaded ODS file '{file_path}' with {len(sheets_data)} sheets")
                return sheets_data
            
        except Exception as e:
            raise ValueError(f"Failed to load ODS file '{file_path}': {e}")

    def _load_numbers_file(self, file_path: str, sheet_name: Optional[str] = None) -> Dict[str, pd.DataFrame]:
        """Load Apple Numbers file into dict of pandas DataFrames (one per sheet/table).
        
        Args:
            file_path: Path to the Numbers file
            sheet_name: Specific sheet to load. If None, load all sheets.
        """
        # Check for Numbers support
        if not NUMBERS_PARSER_AVAILABLE:
            logger.warning("numbers-parser not available - Numbers files cannot be loaded")
            raise ValueError(
                "numbers-parser library is required for .numbers files. "
                "Install with: pip install numbers-parser"
            )
        
        try:
            # Open Numbers document
            doc = Document(file_path)
            logger.info(f"Opened Numbers document with {len(doc.sheets)} sheets")
            
            # Determine which sheets to load
            if sheet_name is not None:
                available_sheet_names = [sheet.name for sheet in doc.sheets]
                if sheet_name not in available_sheet_names:
                    raise ValueError(f"Sheet '{sheet_name}' not found in Numbers file. Available sheets: {available_sheet_names}")
                sheets_to_load = [sheet for sheet in doc.sheets if sheet.name == sheet_name]
                logger.info(f"Loading specific sheet: {sheet_name}")
            else:
                sheets_to_load = doc.sheets
                sheet_names = [sheet.name for sheet in sheets_to_load]
                logger.info(f"Loading all sheets: {sheet_names}")
            
            # Track used table names for uniqueness
            used_names = set()
            sheets_data = {}
            
            for sheet in sheets_to_load:
                logger.info(f"Processing sheet '{sheet.name}' with {len(sheet.tables)} tables")
                
                # Numbers files can have multiple tables per sheet
                for table_idx, table in enumerate(sheet.tables):
                    try:
                        # Get table data as list of lists (includes headers)
                        table_data = table.rows(values_only=True)
                        
                        if not table_data:
                            logger.warning(f"Table {table_idx} in sheet '{sheet.name}' is empty, skipping")
                            continue
                        
                        # First row is typically headers
                        if len(table_data) < 2:
                            logger.warning(f"Table {table_idx} in sheet '{sheet.name}' has no data rows, skipping")
                            continue
                        
                        # Create DataFrame from table data
                        headers = table_data[0]
                        data_rows = table_data[1:]
                        
                        # Handle case where headers might be None or empty
                        if not headers or all(h is None or h == '' for h in headers):
                            # Generate column names
                            headers = [f'Column_{i+1}' for i in range(len(data_rows[0]) if data_rows else 0)]
                        else:
                            # Clean up headers
                            headers = [str(h) if h is not None else f'Column_{i+1}' for i, h in enumerate(headers)]
                        
                        df = pd.DataFrame(data_rows, columns=headers)
                        
                        # Skip empty DataFrames
                        if df.empty:
                            logger.warning(f"Table {table_idx} in sheet '{sheet.name}' resulted in empty DataFrame, skipping")
                            continue
                        
                        # Clean up column names (remove extra whitespace, replace problematic characters)
                        df.columns = [str(col).strip().replace(' ', '_').replace('-', '_').replace('.', '_') for col in df.columns]
                        
                        # Handle Numbers-specific data types and formatting
                        for col in df.columns:
                            if df[col].dtype == 'object':
                                # Try to convert numeric strings to numbers
                                try:
                                    # Check if this looks like a numeric column
                                    numeric_df = pd.to_numeric(df[col], errors='coerce')
                                    non_null_count = numeric_df.notna().sum()
                                    if non_null_count > len(df) * 0.5:  # If >50% can be converted to numbers
                                        df[col] = numeric_df
                                        continue
                                except:
                                    pass
                                
                                # Try to convert to datetime if it looks like dates
                                try:
                                    pd.to_datetime(df[col], errors='raise')
                                    df[col] = pd.to_datetime(df[col], errors='coerce')
                                except:
                                    pass  # Keep as object type
                        
                        # Create unique table name
                        if len(sheet.tables) == 1:
                            # Single table per sheet - use sheet name
                            base_table_name = sheet.name or f"Sheet_{len(sheets_data) + 1}"
                        else:
                            # Multiple tables per sheet - include table index
                            base_table_name = f"{sheet.name}_Table_{table_idx + 1}" if sheet.name else f"Sheet_{len(sheets_data) + 1}_Table_{table_idx + 1}"
                        
                        # Sanitize table name for SQL table name
                        table_name = self._sanitize_sheet_name(base_table_name, used_names)
                        sheets_data[table_name] = df
                        
                        logger.info(f"Successfully loaded table {table_idx} from sheet '{sheet.name}' as table '{table_name}' with {len(df)} rows and {len(df.columns)} columns")
                        
                    except Exception as e:
                        logger.warning(f"Failed to load table {table_idx} from sheet '{sheet.name}': {e}")
                        continue
            
            if not sheets_data:
                raise ValueError(f"Numbers file '{file_path}' contains no readable data in any sheets/tables")
            
            logger.info(f"Successfully loaded Numbers file '{file_path}' with {len(sheets_data)} tables")
            return sheets_data
            
        except Exception as e:
            # Handle specific Numbers parser limitations
            error_msg = str(e).lower()
            if "password" in error_msg or "encrypted" in error_msg:
                raise ValueError(f"Numbers file '{file_path}' is password-protected. Please remove the password and try again.")
            elif "unsupported" in error_msg:
                raise ValueError(f"Numbers file '{file_path}' uses unsupported features. Try re-saving the file in Numbers and try again.")
            else:
                raise ValueError(f"Failed to load Numbers file '{file_path}': {e}")

    def _load_hdf5_file(self, file_path: str, dataset_name: Optional[str] = None) -> Dict[str, pd.DataFrame]:
        """Load HDF5 file into dict of pandas DataFrames (one per dataset).
        
        Args:
            file_path: Path to the HDF5 file
            dataset_name: Specific dataset to load. If None, load all compatible datasets.
        """
        # Check for HDF5 support
        if not H5PY_AVAILABLE:
            logger.warning("h5py not available - HDF5 files cannot be loaded")
            raise ValueError(
                "h5py library is required for .hdf5 files. "
                "Install with: pip install h5py"
            )
        
        try:
            import h5py
            datasets_data = {}
            used_names = set()
            
            with h5py.File(file_path, 'r') as hdf_file:
                logger.info(f"Opened HDF5 file with {len(hdf_file.keys())} top-level items")
                
                # Recursive function to explore HDF5 structure
                def explore_hdf5_group(group, path=""):
                    """Recursively explore HDF5 groups and datasets."""
                    datasets_found = {}
                    
                    for key in group.keys():
                        item = group[key]
                        current_path = f"{path}/{key}" if path else key
                        
                        if isinstance(item, h5py.Dataset):
                            # This is a dataset - try to convert to DataFrame
                            try:
                                # Get dataset info
                                shape = item.shape
                                dtype = item.dtype
                                logger.info(f"Found dataset '{current_path}': shape={shape}, dtype={dtype}")
                                
                                # Skip if dataset_name is specified and doesn't match
                                if dataset_name is not None and key != dataset_name and current_path != dataset_name:
                                    continue
                                
                                # Read dataset
                                data = item[...]
                                
                                # Convert to DataFrame based on shape
                                if len(shape) == 1:
                                    # 1D array - create single column DataFrame
                                    df = pd.DataFrame({key: data})
                                elif len(shape) == 2:
                                    # 2D array - assume rows x columns
                                    if shape[1] == 1:
                                        # Single column
                                        df = pd.DataFrame({key: data.flatten()})
                                    else:
                                        # Multiple columns - generate column names
                                        columns = [f'{key}_col_{i+1}' for i in range(shape[1])]
                                        df = pd.DataFrame(data, columns=columns)
                                else:
                                    # Higher dimensional - flatten to 2D
                                    data_2d = data.reshape(shape[0], -1)
                                    columns = [f'{key}_dim_{i+1}' for i in range(data_2d.shape[1])]
                                    df = pd.DataFrame(data_2d, columns=columns)
                                    logger.info(f"Flattened {len(shape)}D dataset to 2D: {data_2d.shape}")
                                
                                # Handle data types
                                for col in df.columns:
                                    # Try to convert bytes to strings
                                    if df[col].dtype.kind == 'S':  # String/bytes type
                                        try:
                                            df[col] = df[col].astype(str).str.decode('utf-8', errors='ignore')
                                        except:
                                            df[col] = df[col].astype(str)
                                
                                # Create sanitized table name
                                table_name = self._sanitize_sheet_name(current_path.replace('/', '_'), used_names)
                                datasets_found[table_name] = df
                                
                                logger.info(f"Successfully loaded dataset '{current_path}' as table '{table_name}' with {len(df)} rows and {len(df.columns)} columns")
                                
                            except Exception as e:
                                logger.warning(f"Failed to load dataset '{current_path}': {e}")
                                continue
                                
                        elif isinstance(item, h5py.Group):
                            # This is a group - recurse into it
                            logger.info(f"Exploring group '{current_path}' with {len(item.keys())} items")
                            sub_datasets = explore_hdf5_group(item, current_path)
                            datasets_found.update(sub_datasets)
                    
                    return datasets_found
                
                # Start exploration from root
                if dataset_name is not None:
                    # Look for specific dataset
                    available_datasets = []
                    def collect_dataset_names(group, path=""):
                        for key in group.keys():
                            item = group[key]
                            current_path = f"{path}/{key}" if path else key
                            if isinstance(item, h5py.Dataset):
                                available_datasets.append(current_path)
                            elif isinstance(item, h5py.Group):
                                collect_dataset_names(item, current_path)
                    
                    collect_dataset_names(hdf_file)
                    
                    # Check if requested dataset exists
                    if dataset_name not in available_datasets and dataset_name not in [d.split('/')[-1] for d in available_datasets]:
                        raise ValueError(f"Dataset '{dataset_name}' not found in HDF5 file. Available datasets: {available_datasets}")
                
                datasets_data = explore_hdf5_group(hdf_file)
                
                if not datasets_data:
                    raise ValueError(f"HDF5 file '{file_path}' contains no readable datasets")
                
                logger.info(f"Successfully loaded HDF5 file '{file_path}' with {len(datasets_data)} datasets")
                return datasets_data
                
        except Exception as e:
            # Handle specific HDF5 errors
            error_msg = str(e).lower()
            if "permission" in error_msg or "access" in error_msg:
                raise ValueError(f"Cannot access HDF5 file '{file_path}': Permission denied or file in use")
            elif "not an hdf5 file" in error_msg:
                raise ValueError(f"File '{file_path}' is not a valid HDF5 file")
            else:
                raise ValueError(f"Failed to load HDF5 file '{file_path}': {e}")

    def _load_xml_file(self, file_path: str) -> pd.DataFrame:
        """Load XML file into pandas DataFrame."""
        try:
            # Use pandas native XML support (available since pandas 1.3.0)
            df = pd.read_xml(file_path)
            
            # Validate that we have data
            if df.empty:
                raise ValueError(f"XML file '{file_path}' contains no data")
            
            # Clean up column names
            df.columns = [str(col).strip().replace(' ', '_').replace('-', '_') for col in df.columns]
            
            logger.info(f"Successfully loaded XML file '{file_path}' with {len(df)} rows and {len(df.columns)} columns")
            return df
            
        except Exception as e:
            # If pandas XML reading fails, provide helpful error message
            if "lxml" in str(e).lower():
                if not LXML_AVAILABLE:
                    raise ValueError(
                        f"Failed to load XML file '{file_path}': lxml library is required. "
                        "Install with: pip install lxml"
                    )
            raise ValueError(f"Failed to load XML file '{file_path}': {e}")

    def _load_ini_file(self, file_path: str) -> pd.DataFrame:
        """Load INI configuration file into pandas DataFrame."""
        try:
            config = configparser.ConfigParser()
            config.read(file_path)
            
            # Convert INI structure to flat DataFrame format
            rows = []
            for section_name in config.sections():
                section = config[section_name]
                for key, value in section.items():
                    rows.append({
                        'section': section_name,
                        'key': key,
                        'value': value
                    })
            
            # Handle case where no sections exist (only DEFAULT section)
            if not rows and config.defaults():
                for key, value in config.defaults().items():
                    rows.append({
                        'section': 'DEFAULT',
                        'key': key,
                        'value': value
                    })
            
            if not rows:
                raise ValueError(f"INI file '{file_path}' contains no configuration data")
            
            df = pd.DataFrame(rows)
            logger.info(f"Successfully loaded INI file '{file_path}' with {len(df)} configuration entries")
            return df
            
        except Exception as e:
            raise ValueError(f"Failed to load INI file '{file_path}': {e}")

    def _get_table_metadata(self, inspector, table_name):
        columns = inspector.get_columns(table_name)
        foreign_keys = inspector.get_foreign_keys(table_name)
        primary_keys = inspector.get_pk_constraint(table_name)["constrained_columns"]
        indexes = inspector.get_indexes(table_name)
        try:
            table_options = inspector.get_table_options(table_name)
        except NotImplementedError:
            # SQLite and some other dialects don't support table options
            table_options = {}

        col_list = []
        for col in columns:
            col_info = {"name": col["name"], "type": str(col["type"])}
            if col["nullable"] is False:
                col_info["not_null"] = True
            if col.get("autoincrement", False) is True:
                col_info["autoincrement"] = True
            if col.get("default"):
                col_info["default"] = str(col["default"])

            if col["name"] in primary_keys:
                col_info["primary_key"] = True

            for fk in foreign_keys:
                if col["name"] in fk["constrained_columns"]:
                    col_info["foreign_key"] = {
                        "referred_table": fk["referred_table"],
                        "referred_column": fk["referred_columns"][0],
                    }
            col_list.append(col_info)

        index_list = []
        for idx in indexes:
            index_list.append(
                {
                    "name": idx["name"],
                    "columns": idx["column_names"],
                    "unique": idx.get("unique", False),
                }
            )

        return {
            "name": table_name,
            "columns": col_list,
            "foreign_keys": [f["name"] for f in foreign_keys],
            "primary_keys": primary_keys,
            "indexes": index_list,
            "options": table_options,
        }

    # =========================================================
    # Requested Tools
    # =========================================================

    @mcp.tool
    def connect_database(self, name: str, db_type: str, conn_string: str, sheet_name: Optional[str] = None):
        """
        Open a connection to a database.

        Args:
            name: A unique name to identify the connection (e.g., "analytics_db", "user_data").
            db_type: The type of the database ("sqlite", "postgresql", "mysql", "duckdb", "csv", "json", "yaml", "toml", "excel", "ods", "numbers", "xml", "ini", "tsv", "parquet", "feather", "arrow", "hdf5").
            conn_string: The connection string or file path for the database.
            sheet_name: Optional sheet name to load from Excel/ODS/Numbers files, or dataset name for HDF5 files. If not specified, all sheets/datasets are loaded.
        """
        logger.info(f"Attempting to connect to database '{name}' of type '{db_type}'")

        if name in self.connections:
            logger.warning(f"Database '{name}' is already connected")
            return f"Error: A database with the name '{name}' is already connected."

        # Check connection limit
        if not self.connection_semaphore.acquire(blocking=False):
            logger.warning(f"Connection limit reached for database '{name}'")
            return f"Error: Maximum number of concurrent connections (10) reached. Please disconnect a database first."

        try:
            engine = self._get_engine(db_type, conn_string, sheet_name)
            sql_flavor = self._get_sql_flavor(db_type, engine)

            with self.connection_lock:
                self.connections[name] = engine
                self.db_types[name] = db_type
                self.query_history[name] = []
                self.connection_count += 1

            logger.info(
                f"Successfully connected to database '{name}' ({sql_flavor}). Total connections: {self.connection_count}"
            )
            
            response = {
                "success": True,
                "message": f"Successfully connected to database '{name}'",
                "connection_info": {
                    "name": name,
                    "db_type": db_type,
                    "sql_flavor": sql_flavor,
                    "total_connections": self.connection_count
                }
            }
            
            return json.dumps(response, indent=2)
            
        except Exception as e:
            # Release semaphore on failure
            self.connection_semaphore.release()
            logger.error(f"Failed to connect to database '{name}': {e}")
            return f"Failed to connect to database '{name}': {e}"

    @mcp.tool
    def disconnect_database(self, name: str):
        """
        Close a connection to a database. All open connections are closed when the script terminates.

        Args:
            name: The name of the database connection to close.
        """
        logger.info(f"Attempting to disconnect from database '{name}'")
        try:
            conn = self._get_connection(name)
            conn.dispose()

            with self.connection_lock:
                del self.connections[name]
                del self.db_types[name]
                del self.query_history[name]
                self.connection_count -= 1

            # Release semaphore slot
            self.connection_semaphore.release()

            logger.info(
                f"Successfully disconnected from database '{name}'. Total connections: {self.connection_count}"
            )
            return f"Successfully disconnected from database '{name}'."
        except ValueError as e:
            logger.error(f"Database '{name}' not found for disconnection: {e}")
            return str(e)
        except Exception as e:
            logger.error(f"Error disconnecting from database '{name}': {e}")
            return f"An error occurred while disconnecting: {e}"


    @mcp.tool
    def execute_query(self, name: str, query: str, chunk_size: Optional[int] = None) -> str:
        """
        Execute a SQL query and return results as JSON.
        
        For large result sets (>100 rows), automatically creates chunked response with pagination.
        Includes memory usage check before execution to prevent server crashes.
        Auto-clears buffers from the same database when memory is high.

        Args:
            name: The name of the database connection.
            query: The SQL query to execute.
            chunk_size: Optional chunk size for pagination. If not specified, uses default chunking behavior.
        """
        try:
            # Check memory usage before query execution
            memory_info = self._check_memory_usage()
            if memory_info.get("low_memory", False):
                logger.warning(f"High memory usage detected ({memory_info.get('used_percent', 0)}%) before query execution")
                # Auto-clear buffers from same database to free memory
                self._auto_clear_buffers_if_needed(name)

            # Clean up expired buffers
            self._cleanup_expired_buffers()

            engine = self._get_connection(name)
            df = pd.read_sql_query(query, engine)
            self.query_history[name].append(query)

            if df.empty:
                return json.dumps([])

            # Determine chunking threshold - use chunk_size if provided, otherwise 100
            threshold = chunk_size if chunk_size is not None else 100

            # Check row count for large result handling
            if len(df) > threshold:
                # Store full result in buffer
                query_id = self._generate_query_id(name, query)

                # Check if this is a file-based connection to track modifications
                source_file_path = None
                source_file_mtime = None
                connection = self.connections[name]
                if (
                    hasattr(connection, "url")
                    and connection.url.database
                    and connection.url.database != ":memory:"
                ):
                    # This might be a file-based database, but for CSV connections we need to track the original file
                    # For now, we'll leave this as None - more complex file tracking would need connection metadata
                    pass

                buffer = QueryBuffer(
                    query_id=query_id,
                    db_name=name,
                    query=query,
                    results=df,
                    timestamp=time.time(),
                    source_file_path=source_file_path,
                    source_file_mtime=source_file_mtime,
                )

                with self.query_buffer_lock:
                    self.query_buffers[query_id] = buffer

                # Return first chunk with metadata
                chunk_limit = min(chunk_size or 10, len(df))
                first_chunk = df.head(chunk_limit)

                response = {
                    "metadata": {
                        "total_rows": len(df),
                        "showing_rows": f"1-{chunk_limit}",
                        "query_id": query_id,
                        "file_modified_since_buffer": False,  # Will be updated when we implement file tracking
                        "memory_info": memory_info,
                        "chunked": True,
                        "chunk_size": chunk_limit
                    },
                    "data": json.loads(first_chunk.to_json(orient="records")),
                    "pagination": {
                        "use_next_chunk": f"next_chunk(query_id='{query_id}', start_row={chunk_limit + 1}, chunk_size=100)",
                        "get_all_remaining": f"next_chunk(query_id='{query_id}', start_row={chunk_limit + 1}, chunk_size='all')",
                    },
                }

                return json.dumps(response, indent=2)
            else:
                # Return all results for small result sets
                response = {
                    "metadata": {
                        "total_rows": len(df),
                        "showing_rows": f"1-{len(df)}",
                        "memory_info": memory_info,
                        "chunked": False
                    },
                    "data": json.loads(df.to_json(orient="records"))
                }
                return json.dumps(response, indent=2)

        except Exception as e:
            return f"An error occurred while executing the query: {e}"

    @mcp.tool
    def next_chunk(self, query_id: str, start_row: int, chunk_size: str) -> str:
        """
        Retrieve the next chunk of rows from a buffered query result.

        Args:
            query_id: The ID of the buffered query result.
            start_row: The starting row number (1-based indexing).
            chunk_size: Number of rows to retrieve, or 'all' for all remaining rows.
        """
        try:
            # Clean up expired buffers
            self._cleanup_expired_buffers()

            with self.query_buffer_lock:
                if query_id not in self.query_buffers:
                    return f"Error: Query buffer '{query_id}' not found. It may have expired or been cleared."

                buffer = self.query_buffers[query_id]

            df = buffer.results
            total_rows = len(df)

            # Validate start_row (1-based indexing)
            if start_row < 1 or start_row > total_rows:
                return f"Error: start_row must be between 1 and {total_rows}."

            # Convert to 0-based indexing for pandas
            start_idx = start_row - 1

            # Handle chunk_size
            if chunk_size == "all":
                end_idx = total_rows
                chunk_df = df.iloc[start_idx:]
            else:
                try:
                    chunk_size_int = int(chunk_size)
                    if chunk_size_int <= 0:
                        return "Error: chunk_size must be a positive integer or 'all'."
                    end_idx = min(start_idx + chunk_size_int, total_rows)
                    chunk_df = df.iloc[start_idx:end_idx]
                except ValueError:
                    return "Error: chunk_size must be a positive integer or 'all'."

            if chunk_df.empty:
                return json.dumps({"metadata": {"message": "No more rows available"}, "data": []})

            # Build response
            showing_end = start_idx + len(chunk_df)
            response = {
                "metadata": {
                    "query_id": query_id,
                    "total_rows": total_rows,
                    "showing_rows": f"{start_row}-{showing_end}",
                    "chunk_size": len(chunk_df),
                    "buffer_timestamp": buffer.timestamp,
                    "file_modified_since_buffer": self._check_file_modified(buffer),
                },
                "data": json.loads(chunk_df.to_json(orient="records")),
            }

            # Add next pagination options if more rows available
            if showing_end < total_rows:
                next_start = showing_end + 1
                response["pagination"] = {
                    "next_100": f"next_chunk(query_id='{query_id}', start_row={next_start}, chunk_size=100)",
                    "get_all_remaining": f"next_chunk(query_id='{query_id}', start_row={next_start}, chunk_size='all')",
                }

            return json.dumps(response, indent=2)

        except Exception as e:
            return f"An error occurred while retrieving query chunk: {e}"

    @mcp.tool
    def get_query_history(self, name: str) -> str:
        """
        Get the recent query history for a specific database connection.

        Args:
            name: The name of the database connection.
        """
        try:
            history = self.query_history.get(name, [])
            if not history:
                return f"No query history found for database '{name}'."
            return "\n".join(history)
        except Exception as e:
            return f"An error occurred: {e}"

    @mcp.tool
    def list_databases(self) -> str:
        """
        List all available database connections with their SQL flavor information.
        """
        if not self.connections:
            return json.dumps({"message": "No databases are currently connected.", "databases": []})
        
        databases = []
        for name in self.connections.keys():
            db_type = self.db_types.get(name, "unknown")
            sql_flavor = self._get_sql_flavor(db_type, self.connections[name])
            databases.append({
                "name": name,
                "db_type": db_type,
                "sql_flavor": sql_flavor
            })
        
        response = {
            "total_connections": len(databases),
            "databases": databases
        }
        
        return json.dumps(response, indent=2)

    @mcp.tool
    def describe_database(self, name: str) -> str:
        """
        Get detailed information about a database, including its schema in JSON format.

        Args:
            name: The name of the database connection.
        """
        try:
            engine = self._get_connection(name)
            inspector = inspect(engine)

            db_info = {
                "name": name,
                "dialect": engine.dialect.name,
                "version": inspector.get_server_version_info(),
                "default_schema_name": inspector.default_schema_name,
                "schemas": inspector.get_schema_names(),
                "tables": [],
            }

            for table_name in inspector.get_table_names():
                table_info = self._get_table_metadata(inspector, table_name)
                with engine.connect() as conn:
                    safe_table_name = self._safe_table_identifier(table_name)
                    result = conn.execute(
                        text(f"SELECT COUNT(*) FROM {safe_table_name}")  # nosec B608
                    )
                    row_count = result.scalar()
                table_info["size"] = row_count
                db_info["tables"].append(table_info)

            return json.dumps(db_info, indent=2)
        except Exception as e:
            return f"An error occurred: {e}"

    @mcp.tool
    def find_table(self, table_name: str) -> str:
        """
        Find which database contains a specific table.

        Args:
            table_name: The name of the table to find.
        """
        found_dbs = []
        for name, engine in self.connections.items():
            inspector = inspect(engine)
            if table_name in inspector.get_table_names():
                found_dbs.append(name)

        if not found_dbs:
            return f"Table '{table_name}' was not found in any connected databases."
        return json.dumps(found_dbs)

    @mcp.tool
    def describe_table(self, name: str, table_name: str) -> str:
        """
        Get a detailed description of a table including its schema in JSON.

        Args:
            name: The name of the database connection.
            table_name: The name of the table.
        """
        try:
            engine = self._get_connection(name)
            inspector = inspect(engine)
            if table_name not in inspector.get_table_names():
                return (
                    f"Error: Table '{table_name}' does not exist in database '{name}'."
                )

            table_info = self._get_table_metadata(inspector, table_name)
            with engine.connect() as conn:
                safe_table_name = self._safe_table_identifier(table_name)
                result = conn.execute(
                    text(f"SELECT COUNT(*) FROM {safe_table_name}")  # nosec B608
                )
                row_count = result.scalar()
            table_info["size"] = row_count

            return json.dumps(table_info, indent=2)
        except Exception as e:
            return f"An error occurred: {e}"







    def _check_file_modified(self, buffer: QueryBuffer) -> bool:
        """Check if the source file has been modified since buffer creation."""
        if not buffer.source_file_path or not buffer.source_file_mtime:
            return False

        try:
            current_mtime = os.path.getmtime(buffer.source_file_path)
            return current_mtime > buffer.source_file_mtime
        except OSError:
            # File might not exist anymore
            return True

    def _check_memory_usage(self) -> Dict[str, Any]:
        """Check current memory usage and available memory."""
        try:
            memory = psutil.virtual_memory()
            return {
                "total_gb": round(memory.total / (1024**3), 2),
                "available_gb": round(memory.available / (1024**3), 2),
                "used_percent": memory.percent,
                "low_memory": memory.percent > 85  # Consider 85% as low memory threshold
            }
        except Exception as e:
            logger.warning(f"Could not check memory usage: {e}")
            return {"error": str(e), "low_memory": False}

    def _auto_clear_buffers_if_needed(self, db_name: str) -> bool:
        """Auto-clear buffers from the same database if memory is high."""
        memory_info = self._check_memory_usage()
        
        if memory_info.get("low_memory", False):
            with self.query_buffer_lock:
                # Clear buffers from the same database
                buffers_to_clear = [
                    query_id for query_id, buffer in self.query_buffers.items()
                    if buffer.db_name == db_name
                ]
                for query_id in buffers_to_clear:
                    del self.query_buffers[query_id]
                
                if buffers_to_clear:
                    logger.info(f"Auto-cleared {len(buffers_to_clear)} buffers from '{db_name}' due to low memory")
                    return True
        
        return False

    def _get_sql_flavor(self, db_type: str, engine=None) -> str:
        """Determine the SQL flavor/dialect for the database type."""
        if db_type == "sqlite":
            return "SQLite"
        elif db_type == "postgresql":
            return "PostgreSQL"
        elif db_type == "mysql":
            return "MySQL"
        elif db_type == "duckdb":
            return "Duckdb"
        elif db_type == "redis":
            return "Redis"
        elif db_type == "elasticsearch":
            return "Elasticsearch"
        elif db_type == "mongodb":
            return "MongoDB"
        elif db_type == "influxdb":
            return "InfluxDB"
        elif db_type == "neo4j":
            return "Neo4j"
        elif db_type == "couchdb":
            return "CouchDB"
        elif db_type in ["csv", "json", "yaml", "toml", "excel", "ods", "xml", "ini", "tsv", "parquet", "feather", "arrow", "hdf5"]:
            # File formats use SQLite dialect internally
            return "SQLite"
        else:
            # Try to get from engine if available
            if engine and hasattr(engine, 'dialect'):
                return engine.dialect.name.title()
            return "Unknown"

    def _safe_table_identifier(self, table_name: str) -> str:
        """Create a safe SQL identifier for table names to prevent injection."""
        # Validate table name contains only safe characters
        import re

        if not re.match(r"^[a-zA-Z_][a-zA-Z0-9_]*$", table_name):
            raise ValueError(
                f"Invalid table name '{table_name}'. Table names must start with a letter or underscore and contain only alphanumeric characters and underscores."
            )

        # Use SQLAlchemy's quoted_name for safe identifier quoting
        return str(quoted_name(table_name, quote=True))

    # Modern Database Connection Methods
    
    def _create_redis_connection(self, conn_string: str):
        """Create Redis connection from connection string."""
        import redis
        
        # Parse connection string (redis://localhost:6379/0)
        if conn_string.startswith('redis://'):
            # Parse URL format
            parts = conn_string.replace('redis://', '').split('/')
            host_port = parts[0].split(':')
            host = host_port[0] or 'localhost'
            port = int(host_port[1]) if len(host_port) > 1 else 6379
            db = int(parts[1]) if len(parts) > 1 else 0
        else:
            # Default localhost
            host, port, db = 'localhost', 6379, 0
            
        return redis.Redis(host=host, port=port, db=db, decode_responses=True)
    
    def _create_elasticsearch_connection(self, conn_string: str):
        """Create Elasticsearch connection from connection string."""
        from elasticsearch import Elasticsearch
        
        # Parse connection string (http://localhost:9200)
        if not conn_string.startswith('http'):
            conn_string = f"http://{conn_string}"
            
        return Elasticsearch([conn_string])
    
    def _create_mongodb_connection(self, conn_string: str):
        """Create MongoDB connection from connection string."""
        import pymongo
        
        # Parse connection string (mongodb://localhost:27017/database)
        if not conn_string.startswith('mongodb://'):
            conn_string = f"mongodb://{conn_string}"
            
        return pymongo.MongoClient(conn_string)
    
    def _create_influxdb_connection(self, conn_string: str):
        """Create InfluxDB connection from connection string."""
        from influxdb_client import InfluxDBClient
        
        # Parse connection string (http://localhost:8086)
        if not conn_string.startswith('http'):
            conn_string = f"http://{conn_string}"
            
        # InfluxDB requires token and org - use defaults for local testing
        return InfluxDBClient(url=conn_string, token="", org="")
    
    def _create_neo4j_connection(self, conn_string: str):
        """Create Neo4j connection from connection string."""
        from neo4j import GraphDatabase
        
        # Parse connection string (bolt://localhost:7687)
        if not conn_string.startswith('bolt://'):
            conn_string = f"bolt://{conn_string}"
            
        # Neo4j requires auth - use defaults for local testing
        return GraphDatabase.driver(conn_string, auth=("neo4j", "password"))
    
    def _create_couchdb_connection(self, conn_string: str):
        """Create CouchDB connection from connection string."""
        import couchdb
        
        # Parse connection string (http://admin:testpassword@localhost:5984/)
        if not conn_string.startswith('http'):
            conn_string = f"http://admin:testpassword@{conn_string}/"
            
        return couchdb.Server(conn_string)


def main():
    manager = DatabaseManager()
    mcp.run(transport="stdio")


if __name__ == "__main__":
    main()
