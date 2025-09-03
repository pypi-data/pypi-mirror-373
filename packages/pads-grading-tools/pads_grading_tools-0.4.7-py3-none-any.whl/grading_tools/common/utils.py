import pathlib
from collections import defaultdict
from os import PathLike

import openpyxl
import pandas as pd

from grading_tools.common.defaults import *

nd = NamingDictionary()


def complete_file_path(file_path: str, default_file_name: str | None = None, default_extension: str | None = None) -> \
PathLike[str]:
    if default_file_name is None:
        default_file_name = 'output'
    if default_extension is None:
        default_extension = '.xlsx'

    path = pathlib.Path(file_path)
    if path.is_dir():
        if not default_file_name.endswith('.csv') and not default_file_name.endswith('.xlsx'):
            default_file_name += default_extension
        path = path.joinpath(default_file_name)
    return path


def get_possible_grades(numbers: bool = True, absent: bool = True, symbols: bool = False, decimal_sep: str = '.'):
    possible_grades = []
    if numbers:
        for x in range(1, 4):
            possible_grades.extend((f'{x}{decimal_sep}0', f'{x}{decimal_sep}3', f'{x}{decimal_sep}7'))
    possible_grades.append(f'4{decimal_sep}0')
    possible_grades.append(f'5{decimal_sep}0')
    if absent or symbols:
        possible_grades.append('X')
    if symbols:
        possible_grades.extend(('NZ', 'PA', 'U', 'Q-Q'))
    return possible_grades


def read_excel_table(file_name: str, table_name: str) -> pd.DataFrame:
    wb = None
    try:
        wb = openpyxl.load_workbook(file_name, read_only=False,
                                    data_only=True)  # openpyxl does not have table info if read_only is True; data_only means any functions will pull the last saved value instead of the formula
        sheet, tbl_range = None, None
        for sheetname in wb.sheetnames:  # pulls as strings
            sheet = wb[sheetname]  # get the sheet object instead of string
            if table_name in sheet.tables:  # tables are stored within sheets, not within the workbook, although table names are unique in a workbook
                tbl = sheet.tables[table_name]  # get table object instead of string
                tbl_range = tbl.ref  # something like 'C4:F9'
                break  # we've got our table, bail from for-loop
        if tbl_range is None:
            print(f'Table {table_name} not found')
            return None
        data = sheet[tbl_range]  # returns a tuple that contains rows, where each row is a tuple containing cells
        content = [[cell.value for cell in row] for row in data]  # loop through those row/cell tuples
        header = content[0]  # first row is column headers
        rest = content[1:]  # every row that isn't the first is data
        df = pd.DataFrame(rest, columns=header)
        return df
    finally:
        if wb:
            wb.close()


def read_generic_table(file_name: str, columns: list[str] | None = None, sheet_name: str | None = None,
                       table_name: str | None = None, **kwargs) -> pd.DataFrame:
    print(f'Reading {table_name}/{sheet_name} from to {file_name}.')

    df = None
    if file_name.endswith('.csv'):
        df = pd.read_csv(file_name)
    elif file_name.endswith('.xlsx'):
        if table_name is not None:
            df = read_excel_table(file_name, table_name)
        elif sheet_name is not None:
            df = pd.read_excel(file_name, sheet_name)

    if columns is not None:
        df = df[columns]

    return df


def read_rwthonline(file_name: str) -> pd.DataFrame:
    col_dtypes = defaultdict(lambda: 'str')
    col_dtypes |= {'REGISTRATION_NUMBER': 'Int64'}
    return pd.read_csv(file_name, sep=nd.RWTHONLINE_CSV_SEP, quotechar=nd.RWTHONLINE_QUOTE_CHAR, dtype=col_dtypes)


def read_moodle_csv(file_name: str | PathLike[str]) -> pd.DataFrame:
    return pd.read_csv(file_name, sep=nd.MOODLE_CSV_SEP, quotechar=nd.MOODLE_QUOTE_CHAR,
                       decimal=nd.MOODLE_CSV_DECIMAL_SEP, na_values=[nd.MOODLE_NA_CHAR])


def read_moodle_students_csv(file_name: str | PathLike[str]) -> pd.DataFrame:
    return pd.read_csv(file_name, sep=nd.MOODLE_GROUPS_CSV_SEP, quotechar=nd.MOODLE_GROUPS_QUOTE_CHAR,
                       decimal=nd.MOODLE_CSV_DECIMAL_SEP, na_values=[nd.MOODLE_NA_CHAR])


# for custom grading aspects that don't belong to a moodle activity, e.g., exam scores
def read_moodle_excel(file_name: str | PathLike[str]) -> pd.DataFrame:
    return pd.read_excel(file_name, na_values=[nd.MOODLE_NA_CHAR])


def write_moodle_csv(df: pd.DataFrame, file_path: str | PathLike[str]):
    p = pathlib.Path(file_path)
    print(f'Writing Moodle Upload to {p}.')
    p.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(file_path, index=False, sep=nd.MOODLE_CSV_SEP, quotechar=nd.MOODLE_QUOTE_CHAR,
              decimal=nd.MOODLE_CSV_DECIMAL_SEP,
              encoding='utf-8')


def write_rwthonline_csv(df: pd.DataFrame, file_path: str | PathLike[str]):
    p = pathlib.Path(file_path).resolve()
    print(f'Writing RWTH Online CSV to {p}.')
    p.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(file_path, index=False, sep=nd.RWTHONLINE_CSV_SEP, quotechar=nd.RWTHONLINE_QUOTE_CHAR, encoding='utf-8')
